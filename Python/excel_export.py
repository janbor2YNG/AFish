"""Excel export of a campaign's results - one of the most important features:
a plain .xlsx file with three sheets (Uebersicht, Teilnehmeruebersicht,
Detaildaten) that teachers/evaluators can open without any extra tooling.
"""

import io
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(BASE_DIR, "..", "Server", "Flask"))

import stats as stats_mod  # noqa: E402

BOLD = Font(bold=True)


def _autosize(ws):
    for col in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(length + 2, 10), 40)


def build_campaign_workbook(campaign_id) -> io.BytesIO:
    data = stats_mod.campaign_stats(campaign_id)
    if data is None:
        raise FileNotFoundError(f"campaign {campaign_id} not found")

    wb = Workbook()

    # --- Sheet 1: Uebersicht --------------------------------------------------
    ws = wb.active
    ws.title = "Uebersicht"
    ws.append(["Kennzahl", "Wert"])
    for cell in ws[1]:
        cell.font = BOLD
    by_cat = data["totals"]["by_category"]
    rows = [
        ("Kampagne", data["meta"].get("name")),
        ("Zeitraum", f"{data['meta'].get('start_date')} - {data['meta'].get('end_date')}"),
        ("Gesamt versendete Mails", data["totals"]["sent"]),
    ]
    for cat, count in sorted(by_cat.items()):
        rows.append((f"{cat.capitalize()}-Mails", count))
    rows += [
        ("Klicks (nicht erkannt)", data["totals"]["clicked"]),
        ("Erfolgreich gemeldet", data["totals"]["reported"]),
        ("Klickquote (%)", data["click_rate"]),
        ("Erkennungsquote (%)", data["report_rate"]),
    ]
    for r in rows:
        ws.append(r)
    _autosize(ws)

    # --- Sheet 2: Teilnehmeruebersicht -----------------------------------------
    ws2 = wb.create_sheet("Teilnehmeruebersicht")
    header = ["Name", "E-Mail", "Gruppe", "E-Mails erhalten", "Geklickt",
               "Korrekt gemeldet", "Mehrfach nicht erkannt"]
    ws2.append(header)
    for cell in ws2[1]:
        cell.font = BOLD
    for u in data["users"]:
        ws2.append([
            u["name"], u["email"], u["gruppe"] or "", u["sent"], u["clicked"],
            u["reported"], "Ja" if u["repeat_fail"] else "Nein",
        ])
    _autosize(ws2)

    # --- Sheet 3: Detaildaten ---------------------------------------------------
    ws3 = wb.create_sheet("Detaildaten")
    header3 = ["Empfaenger", "E-Mail", "Welle", "Kategorie", "Template",
               "Absender", "Status", "Versendet am", "Geklickt am", "Gemeldet am"]
    ws3.append(header3)
    for cell in ws3[1]:
        cell.font = BOLD
    status_labels = {"sent": "zugestellt", "clicked": "geklickt", "reported": "gemeldet"}
    for e in data["events"]:
        ws3.append([
            e["user_name"], e["user_email"], e["wave"], e["category"], e["template"],
            f"{e['sender_name']} <{e['sender_email']}>",
            status_labels.get(e["status"], e["status"]),
            e["sent_at"] or "", e["clicked_at"] or "", e["reported_at"] or "",
        ])
    _autosize(ws3)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
